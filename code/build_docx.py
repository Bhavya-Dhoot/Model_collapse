"""Build paper/main.docx in IEEE conference layout from the LaTeX sources.

No TeX toolchain is needed: python-docx lays out the page, latex2mathml plus
Word's own MML2OMML.XSL turn every $...$ / display into native (editable) Word
equations, and refs.bib is rendered in IEEE citation order. Content order
follows main.tex exactly; floats are placed at their source position.
"""
import os
import re

import bibtexparser
import latex2mathml.converter as l2m
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt, Twips
from lxml import etree

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "paper")
OUT = os.path.join(ROOT, "main.docx")
XSL = r"C:\Program Files (x86)\Microsoft Office\root\Office16\MML2OMML.XSL"
COL_W = Inches(3.5)
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
WARN = []

# ----------------------------------------------------------------- LaTeX text
MACROS = [(r"\alphastar", r"\alpha^{\star}"), (r"\real", r"\mathcal{D}_{\mathrm{r}}"),
          (r"\hold", r"\mathcal{H}"), (r"\Preal", "P"), (r"\E", r"\mathbb{E}"),
          (r"\Var", r"\mathrm{Var}"), (r"\Cov", r"\mathrm{Cov}"), (r"\Sig", r"\Sigma"),
          (r"\W", r"\mathcal{W}")]


def strip_comments(s):
    return re.sub(r"(?m)(?<!\\)%.*$", "", s)


def norm_ws(s):
    return re.sub(r"\s+", " ", s).strip()


def read_group(s, i):
    """s[i] == '{' -> (content, index after the matching '}')"""
    assert s[i] == "{", s[i:i + 20]
    depth, j = 0, i
    while j < len(s):
        if s[j] == "\\":
            j += 2
            continue
        if s[j] == "{":
            depth += 1
        elif s[j] == "}":
            depth -= 1
            if depth == 0:
                return s[i + 1:j], j + 1
        j += 1
    raise ValueError("unbalanced braces near " + s[i:i + 40])


def find_arg(s, cmd):
    """(content, start, end) of the first \\cmd{...} in s, brace matched; or None"""
    m = re.search(r"\\" + cmd + r"\s*\{", s)
    if not m:
        return None
    content, end = read_group(s, m.end() - 1)
    return content, m.start(), end


def expand_macros(s):
    for k, v in MACROS:
        s = re.sub(re.escape(k) + r"(?![A-Za-z])", lambda m, v=v: v, s)
    return s


def prep_math(s):
    s = expand_macros(s)
    s = re.sub(r"\\allowbreak", "", s)
    s = re.sub(r"\\(?:bigl|bigr|Bigl|Bigr|big|Big)\b", "", s)
    s = re.sub(r"\\[td]frac", r"\\frac", s)
    s = re.sub(r"\\limsup", r"\\mathrm{lim}\\,\\mathrm{sup}", s)
    s = re.sub(r"\\(hat|bar|tilde|mathrm|mathcal|mathbb)\s+(\\?[A-Za-z]+)", r"\\\1{\2}", s)
    s = re.sub(r"\\(hat|bar|tilde)(\\[A-Za-z]+)", r"\\\1{\2}", s)
    s = re.sub(r"([_^])\\([A-Za-z]+)", r"\1{\\\2}", s)
    s = re.sub(r"\{(=|\\to|\\approx|\\le|\\ge|\\sim)\}", r"\1", s)
    # \underbrace{X}_{label} -> label set directly under X (Word has no clean brace-with-label)
    s = re.sub(r"\\underbrace\{([^{}]*)\}_\{((?:[^{}]|\{[^{}]*\})*)\}", r"\\underset{\2}{\1}", s)
    return s.strip()


# latex2mathml emits \hat/\bar/\tilde as plain <mover> with ASCII ^ ~ ¯, which Word's XSL turns
# into floating limits; mark them as true accents with combining characters so Word uses m:acc.
ACCENT_MAP = {"&#x0005E;": "&#x00302;", "&#x0007E;": "&#x00303;", "&#x000AF;": "&#x00304;"}


def fix_accents(mml):
    def repl(m):
        return f'<mover accent="true">{m.group(1)}<mo>{ACCENT_MAP[m.group(2)]}</mo></mover>'
    return re.sub(r'<mover>(<mrow>.*?</mrow>)<mo stretchy="(?:true|false)">(&#x0005E;|&#x0007E;|&#x000AF;)</mo></mover>',
                  repl, mml)


XSLT = etree.XSLT(etree.parse(XSL))


def omml(latex, display=False):
    """LaTeX math -> OMML element: m:oMath (inline) or m:oMathPara (display)."""
    src = prep_math(latex)
    try:
        mml = fix_accents(l2m.convert(src, display="block" if display else "inline"))
        root = XSLT(etree.fromstring(mml.encode())).getroot()
    except Exception as e:  # noqa: BLE001
        WARN.append(f"math failed: {latex!r}: {e}")
        return None
    is_para = root.tag == f"{{{M_NS}}}oMathPara"
    if display:
        if is_para:
            return root
        para = etree.Element(f"{{{M_NS}}}oMathPara")
        para.append(root)
        return para
    return root.find(f"{{{M_NS}}}oMath") if is_para else root


# ----------------------------------------------------------------- numbering
def roman(n):
    out = ""
    for v, r in [(10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")]:
        while n >= v:
            out += r
            n -= v
    return out


THM_ENVS = ("proposition", "definition", "remark", "theorem", "lemma", "corollary")


class Numbering:
    """Citation order and every \\label -> printed number, mirroring IEEEtran."""

    def __init__(self, body):
        self.cites, self.labels = [], {}
        for m in re.finditer(r"\\cite\{([^}]*)\}", body):
            for k in m.group(1).split(","):
                k = k.strip()
                if k not in self.cites:
                    self.cites.append(k)
        sec = sub = fig = tab = eq = 0
        env_ct = {e: 0 for e in THM_ENVS}
        pat = re.compile(r"\\section\*?\{|\\subsection\{|\\begin\{(figure\*?|table|equation|"
                         + "|".join(THM_ENVS) + r")\}")
        pos = 0
        while True:
            m = pat.search(body, pos)
            if not m:
                break
            tok = m.group(0)
            if tok.startswith("\\section*"):
                pos = m.end()
                continue
            if tok.startswith("\\section") or tok.startswith("\\subsection"):
                if tok.startswith("\\section"):
                    sec += 1
                    sub = 0
                    num = roman(sec)
                else:
                    sub += 1
                    num = f"{roman(sec)}-{chr(64 + sub)}"
                nxt = pat.search(body, m.end())
                block = body[m.end():nxt.start() if nxt else len(body)]
                for lb in re.findall(r"\\label\{(sec:[^}]*)\}", block):
                    self.labels.setdefault(lb, num)
                pos = m.end()
                continue
            env = m.group(1)
            endm = re.search(r"\\end\{" + re.escape(env) + r"\}", body[m.end():])
            block = body[m.end():m.end() + endm.start()]
            if env.startswith("figure"):
                fig += 1
                num = str(fig)
            elif env == "table":
                tab += 1
                num = roman(tab)
            elif env == "equation":
                eq += 1
                num = str(eq)
            else:
                env_ct[env] += 1
                num = str(env_ct[env])
            for lb in re.findall(r"\\label\{([^}]*)\}", block):
                if not lb.startswith("sec:"):
                    self.labels.setdefault(lb, num)
            pos = m.end()

    def cite(self, keys):
        nums = sorted(self.cites.index(k.strip()) + 1 for k in keys.split(","))
        out, i = [], 0
        while i < len(nums):
            j = i
            while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
                j += 1
            if j - i >= 2:
                out.append(f"[{nums[i]}]\u2013[{nums[j]}]")
            else:
                out.append(", ".join(f"[{n}]" for n in nums[i:j + 1]))
            i = j + 1
        return ", ".join(out)

    def ref(self, key):
        if key not in self.labels:
            WARN.append(f"unresolved ref {key}")
            return "??"
        return self.labels[key]


# ----------------------------------------------------------------- docx helpers
def set_cols(section, n, space_twips=360):
    sectPr = section._sectPr
    cols = sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sectPr.append(cols)
    cols.set(qn("w:num"), str(n))
    cols.set(qn("w:space"), str(space_twips))
    cols.set(qn("w:equalWidth"), "1")


def page_setup(section):
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(0.625)
    section.top_margin, section.bottom_margin = Inches(0.75), Inches(1.0)


def new_para(doc, style=None, align=None, indent=True, size=None, before=0, after=0, keep_next=False):
    p = doc.add_paragraph(style=style)
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent = Inches(0.14) if indent else None
    pf.keep_with_next = keep_next
    p._size = size
    return p


def run(p, text, bold=None, italic=None, size=None, small_caps=None, mono=False):
    r = p.add_run(text)
    if bold is not None:
        r.bold = bold
    if italic is not None:
        r.italic = italic
    sz = size or getattr(p, "_size", None)
    if sz:
        r.font.size = Pt(sz)
    if small_caps:
        r.font.small_caps = True
    if mono:
        r.font.name = "Courier New"
        r._element.rPr.rFonts.set(qn("w:hAnsi"), "Courier New")
    return r


def add_tab(p, pos, align=WD_TAB_ALIGNMENT.RIGHT):
    p.paragraph_format.tab_stops.add_tab_stop(pos, align)


def set_cell_borders(cell, **edges):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders")
        tcPr.append(b)
    for edge, val in edges.items():
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single" if val else "nil")
        e.set(qn("w:sz"), str(val or 0))
        e.set(qn("w:color"), "000000")
        b.append(e)


# ----------------------------------------------------------------- inline text
TEXT_REPL = [("---", "\u2014"), ("--", "\u2013"), ("``", "\u201c"), ("''", "\u201d"), ("~", "\u00a0"),
             (r"\%", "%"), (r"\&", "&"), (r"\,", "\u2009"), (r"\ ", " "), (r"\_", "_"), (r"\#", "#"),
             (r"\ldots", "\u2026")]
STYLE_CMDS = ("emph", "textit", "textbf", "texttt", "textsc")
NOOP_CMDS = ("centering", "small", "footnotesize", "noindent", "allowbreak")


def emit_inline(p, s, N, italic=False, bold=False):
    """Walk LaTeX inline text, emitting runs / OMML into paragraph p."""
    i, buf = 0, ""

    def flush():
        nonlocal buf
        if buf:
            t = buf
            for a, b in TEXT_REPL:
                t = t.replace(a, b)
            run(p, t, bold=bold or None, italic=italic or None)
            buf = ""

    while i < len(s):
        c = s[i]
        if c == "$":
            j = s.find("$", i + 1)
            inner = s[i + 1:j]
            if re.fullmatch(r"[\d.\s/+\-]+", inner.replace("\\,", "")):  # plain numbers: keep the text font
                buf += inner.replace("\\,", "")
                i = j + 1
                continue
            flush()
            el = omml(inner)
            if el is not None:
                sz = getattr(p, "_size", None)
                if sz:  # match the paragraph's font size (captions, tables, abstract)
                    for r in el.iter(f"{{{M_NS}}}r"):
                        rpr = OxmlElement("w:rPr")
                        szel = OxmlElement("w:sz")
                        szel.set(qn("w:val"), str(int(sz * 2)))
                        rpr.append(szel)
                        mrpr = r.find(f"{{{M_NS}}}rPr")
                        r.insert(list(r).index(mrpr) + 1 if mrpr is not None else 0, rpr)
                p._p.append(el)
            else:
                run(p, s[i + 1:j], italic=True)
            i = j + 1
            continue
        if c == "\\":
            if s.startswith("\\\\", i):
                flush()
                p.add_run().add_break(WD_BREAK.LINE)
                i += 2
                m2 = re.match(r"\[[^\]]*\]", s[i:])
                i += m2.end() if m2 else 0
                continue
            m = re.match(r"\\([A-Za-z]+)\*?", s[i:])
            if not m:  # escaped symbol such as \% or \,
                buf += s[i:i + 2]
                i += 2
                continue
            cmd, j = m.group(1), i + m.end()
            while j < len(s) and s[j] == " " and cmd in ("cite", "ref", "eqref", "label"):
                j += 1
            if cmd in STYLE_CMDS:
                arg, k = read_group(s, j)
                flush()
                if cmd == "texttt":
                    run(p, arg, mono=True, italic=italic or None)
                elif cmd == "textbf":
                    emit_inline(p, arg, N, italic=italic, bold=True)
                else:
                    emit_inline(p, arg, N, italic=not italic, bold=bold)
                i = k
                continue
            if cmd == "texorpdfstring":
                a, k = read_group(s, j)
                _, k = read_group(s, k)
                flush()
                emit_inline(p, a, N, italic=italic, bold=bold)
                i = k
                continue
            if cmd == "cite":
                arg, k = read_group(s, j)
                buf += N.cite(arg)
                i = k
                continue
            if cmd == "ref":
                arg, k = read_group(s, j)
                buf += N.ref(arg)
                i = k
                continue
            if cmd == "eqref":
                arg, k = read_group(s, j)
                buf += "(" + N.ref(arg) + ")"
                i = k
                continue
            if cmd == "label":
                _, k = read_group(s, j)
                i = k
                continue
            if cmd in NOOP_CMDS:
                i = j
                continue
            if cmd == "setlength":
                _, k = read_group(s, j)
                _, k = read_group(s, k)
                i = k
                continue
            WARN.append(f"unknown inline command \\{cmd}")
            buf += "\\" + cmd
            i = j
            continue
        buf += c
        i += 1
    flush()


def split_displays(text):
    """-> [('t', text) | ('d', latex)] chunks, splitting out \\[ ... \\]"""
    pos, chunks = 0, []
    for m in re.finditer(r"\\\[(.*?)\\\]|\\begin\{(gather\*|align\*?)\}(.*?)\\end\{\2\}", text, re.S):
        chunks.append(("t", text[pos:m.start()]))
        if m.group(1) is not None:
            chunks.append(("d", m.group(1)))
        else:  # one display per line of a gather*/align* block
            for line in re.split(r"\\\\", m.group(3)):
                if line.strip():
                    chunks.append(("d", line))
        pos = m.end()
    chunks.append(("t", text[pos:]))
    return chunks


# ----------------------------------------------------------------- block emitters
class Builder:
    def __init__(self, doc, N):
        self.doc, self.N = doc, N
        self.sec = self.sub = 0
        self.env_ct = {e: 0 for e in THM_ENVS}

    def paragraphs(self, text, italic=False, size=None):
        for kind, c in split_displays(text):
            if kind == "d":
                self.display(c)
                continue
            for part in re.split(r"\n\s*\n", c):
                part = norm_ws(part)
                if part:
                    self.para(part, italic, True, size)

    def para(self, text, italic=False, indent=True, size=None):
        if not re.sub(r"\\label\{[^}]*\}", "", text).strip():  # a lone \label is not a paragraph
            return None
        p = new_para(self.doc, indent=indent, size=size)
        emit_inline(p, text, self.N, italic=italic)
        return p

    def display(self, latex, number=None):
        p = new_para(self.doc, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, before=3, after=3)
        el = omml(latex, display=True)
        if number:
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_tab(p, Twips(int(COL_W.twips / 2)), WD_TAB_ALIGNMENT.CENTER)
            add_tab(p, COL_W, WD_TAB_ALIGNMENT.RIGHT)
            run(p, "\t")
        if el is None:
            run(p, latex, italic=True)
        elif number:
            p._p.append(el.find(f"{{{M_NS}}}oMath"))
        else:
            p._p.append(el)
        if number:
            run(p, f"\t({number})")

    def heading(self, title, num):
        p = new_para(self.doc, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, before=8, after=4, keep_next=True)
        if num:
            run(p, f"{num}. ", size=10)
        run(p, title, size=10, small_caps=True)

    def subheading(self, title, num):
        p = new_para(self.doc, align=WD_ALIGN_PARAGRAPH.LEFT, indent=False, before=6, after=3, keep_next=True)
        run(p, f"{num}. ", italic=True, size=10)
        emit_inline(p, title, self.N, italic=True)

    def theorem(self, env, block):
        name = env.capitalize()
        title = None
        m = re.match(r"\s*\[", block)
        if m:
            j = block.find("]")
            title, block = block[m.end():j], block[j + 1:]
        self.env_ct[env] += 1  # same per-environment counters IEEEtran uses; labels are optional
        num = str(self.env_ct[env])
        body_italic = env in ("proposition", "theorem", "lemma", "corollary")
        head = f"{name} {num}" + (f" ({norm_ws(title)})" if title else "") + ": "
        started = False
        for kind, c in split_displays(block.strip()):
            if kind == "d":
                self.display(c)
                continue
            for part in re.split(r"\n\s*\n", c):
                part = norm_ws(part)
                if not part:
                    continue
                p = new_para(self.doc, indent=True, before=0 if started else 3)
                if not started:
                    run(p, head, italic=True)
                    started = True
                emit_inline(p, part, self.N, italic=body_italic)

    def proof(self, block):
        label = "Proof"
        m = re.match(r"\s*\[([^\]]*)\]", block)  # \begin{proof}[Proof sketch]
        if m:
            label, block = m.group(1), block[m.end():]
        started, last = False, None
        for kind, c in split_displays(block.strip()):
            if kind == "d":
                self.display(c)
                last = None
                continue
            for part in re.split(r"\n\s*\n", c):
                part = norm_ws(part)
                if not part:
                    continue
                p = new_para(self.doc, indent=True)
                if not started:
                    run(p, f"{label}: ", italic=True)
                    started = True
                emit_inline(p, part, self.N)
                last = p
        if last is None:
            last = new_para(self.doc, indent=False)
        add_tab(last, COL_W, WD_TAB_ALIGNMENT.RIGHT)
        run(last, "\t\u220e")

    def itemize(self, block):
        for item in re.split(r"\\item\b", block)[1:]:
            p = self.doc.add_paragraph(style="List Bullet")
            pf = p.paragraph_format
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf.space_after = Pt(0)
            pf.left_indent = Inches(0.2)
            emit_inline(p, norm_ws(item), self.N)

    def figure(self, block, star):
        cap = find_arg(block, "caption")[0]
        lb = re.search(r"\\label\{([^}]*)\}", block)
        num = self.N.labels.get(lb.group(1), "?") if lb else "?"
        g = re.search(r"\\includegraphics(?:\[[^\]]*\])?\{([^}]*)\}", block)
        img = os.path.join(ROOT, g.group(1).replace(".pdf", ".png")) if g else os.path.join(ROOT, "figs", "fig_loop.png")
        if star:  # full-width float: single-column continuous section
            s = self.doc.add_section(WD_SECTION.CONTINUOUS)
            page_setup(s)
            set_cols(s, 1)
        p = new_para(self.doc, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, before=4, after=2, keep_next=True)
        p.add_run().add_picture(img, width=Inches(7.0) if star else COL_W)
        c = new_para(self.doc, align=WD_ALIGN_PARAGRAPH.CENTER if star else WD_ALIGN_PARAGRAPH.JUSTIFY,
                     indent=False, after=6, size=8)
        run(c, f"Fig. {num}. ", size=8)
        emit_inline(c, norm_ws(cap), self.N)
        if star:
            s = self.doc.add_section(WD_SECTION.CONTINUOUS)
            page_setup(s)
            set_cols(s, 2)

    def table(self, block):
        cap, cs, ce = find_arg(block, "caption")
        lb = re.search(r"\\label\{([^}]*)\}", block)
        num = self.N.labels.get(lb.group(1), "?") if lb else "?"
        p = new_para(self.doc, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, before=4, size=8, keep_next=True)
        run(p, f"TABLE {num}", size=8)
        c = new_para(self.doc, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=False, after=3, size=8)
        emit_inline(c, norm_ws(cap), self.N)
        body = block[:cs] + block[ce:]
        body = re.sub(r"\\label\{[^}]*\}|\\centering|\\small|\\setlength\{[^}]*\}\{[^}]*\}", "", body)
        body = re.sub(r"^\s*\[[^\]]*\]", "", body)  # float placement option, e.g. [t]
        pos = 0
        while True:
            m = re.search(r"\\begin\{tabular\}", body[pos:])
            if not m:
                break
            start = pos + m.start()
            spec, k = read_group(body, pos + m.end())  # brace-matched: handles @{} in the spec
            endm = re.search(r"\\end\{tabular\}", body[k:])
            pre = norm_ws(re.sub(r"\\\\\[[^\]]*\]", "", body[pos:start]))
            if pre:
                q = new_para(self.doc, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, size=8, keep_next=True)
                emit_inline(q, pre, self.N)
            self.tabular(spec, body[k:k + endm.start()])
            pos = k + endm.end()
        gap = new_para(self.doc, indent=False, after=2)  # small spacer; a 10pt empty line is too tall
        rpr = OxmlElement("w:rPr")
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), "8")
        rpr.append(sz)
        gap._p.get_or_add_pPr().append(rpr)

    def tabular(self, spec, content):
        ncol = len(re.sub(r"@\{[^}]*\}|\|", "", spec))
        rows, rules = [], []
        for line in content.strip().split("\n"):
            line = line.strip()
            if not line or line.startswith("\\toprule") or line.startswith("\\bottomrule"):
                continue
            if line.startswith("\\midrule"):
                rules.append(len(rows))
                continue
            rows.append([c.strip() for c in re.split(r"(?<!\\)&", line.rstrip("\\").rstrip())])
        t = self.doc.add_table(rows=len(rows), cols=ncol)
        t.alignment = 1
        t.autofit = False
        # column widths proportional to the longest *body* cell as it will print (macros stripped),
        # floored so short labels still fit; headers may wrap, body values must not
        def vis_len(s):
            s = re.sub(r"\\text\{([^}]*)\}", r"\1", s)
            s = re.sub(r"\\[A-Za-z]+", "", s)
            return len(re.sub(r"[{}$_^\\,;]", "", s))
        body_rows = rows[rules[0]:] if rules else rows
        lens = [max([vis_len(r[ci]) if ci < len(r) else 0 for r in body_rows] + [7]) for ci in range(ncol)]
        widths = [Emu(int(COL_W * ln / sum(lens))) for ln in lens]
        for ci in range(ncol):
            t.columns[ci].width = widths[ci]
        # tight cell margins (IEEE tables are dense) so 8pt numbers never wrap
        tblPr = t._tbl.tblPr
        mar = OxmlElement("w:tblCellMar")
        for side in ("left", "right"):
            e = OxmlElement(f"w:{side}")
            e.set(qn("w:w"), "40")
            e.set(qn("w:type"), "dxa")
            mar.append(e)
        tblPr.append(mar)
        for ri, cells in enumerate(rows):
            trPr = t.rows[ri]._tr.get_or_add_trPr()
            trPr.append(OxmlElement("w:cantSplit"))
            for ci in range(ncol):
                cell = t.cell(ri, ci)
                cell.width = widths[ci]
                cp = cell.paragraphs[0]
                cp.paragraph_format.space_after = Pt(0)
                cp.paragraph_format.keep_with_next = ri < len(rows) - 1  # keep the table in one column
                cp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER
                cp._size = 8
                if ci < len(cells):
                    emit_inline(cp, cells[ci], self.N)
                top = 8 if ri == 0 or ri in rules else 0
                bottom = 8 if ri == len(rows) - 1 else 0
                set_cell_borders(cell, top=top, bottom=bottom, left=0, right=0)

    def body(self, text):
        """Top-level walk of the LaTeX body, in source order."""
        pat = re.compile(r"\\section\*?\{|\\subsection\{|\\begin\{(figure\*?|table|itemize|proof|gather\*|equation|align\*?|"
                         + "|".join(THM_ENVS) + r")\}|\\\[")
        pos = 0
        while True:
            m = pat.search(text, pos)
            if not m:
                self.paragraphs(text[pos:])
                break
            self.paragraphs(text[pos:m.start()])
            tok = m.group(0)
            if tok.startswith("\\section") or tok.startswith("\\subsection"):
                title, k = read_group(text, m.end() - 1)
                if tok.startswith("\\section*"):
                    self.heading(norm_ws(title), None)
                elif tok.startswith("\\section"):
                    self.sec += 1
                    self.sub = 0
                    self.heading(norm_ws(title), roman(self.sec))
                else:
                    self.sub += 1
                    self.subheading(norm_ws(title), chr(64 + self.sub))
                m2 = re.match(r"\s*\\label\{[^}]*\}", text[k:])
                pos = k + (m2.end() if m2 else 0)
                continue
            if tok == "\\[":
                j = text.find("\\]", m.end())
                self.display(text[m.end():j])
                pos = j + 2
                continue
            env = m.group(1)
            endm = re.search(r"\\end\{" + re.escape(env) + r"\}", text[m.end():])
            block = text[m.end():m.end() + endm.start()]
            pos = m.end() + endm.end()
            if env.startswith("figure"):
                self.figure(block, env.endswith("*"))
            elif env == "table":
                self.table(block)
            elif env == "itemize":
                self.itemize(block)
            elif env == "proof":
                self.proof(block)
            elif env in ("gather*", "align*", "align"):
                for line in re.split(r"\\\\", block):
                    if line.strip():
                        self.display(line)
            elif env == "equation":
                lb = re.search(r"\\label\{([^}]*)\}", block)
                num = self.N.labels.get(lb.group(1)) if lb else None
                self.display(re.sub(r"\\label\{[^}]*\}", "", block), number=num)
            else:
                self.theorem(env, block)


# ----------------------------------------------------------------- references
ACCENTS = {"'o": "\u00f3", "'a": "\u00e1", "'e": "\u00e9", "'i": "\u00ed", '"a': "\u00e4", '"o': "\u00f6", '"u': "\u00fc"}


def tex_to_text(s):
    s = re.sub(r"\{\\(['\"])([a-zA-Z])\}", lambda m: ACCENTS.get(m.group(1) + m.group(2), m.group(2)), s)
    return norm_ws(s.replace("{", "").replace("}", "").replace("--", "\u2013"))


def fmt_authors(a):
    names = []
    for person in re.split(r"\s+and\s+", a):
        person = tex_to_text(person)
        if "," in person:
            last, first = [x.strip() for x in person.split(",", 1)]
        else:
            parts = person.split()
            last, first = parts[-1], " ".join(parts[:-1])
        initials = " ".join("-".join(f"{x[0]}." for x in w.split("-") if x) for w in re.split(r"[\s.]+", first) if w)
        names.append(f"{initials} {last}".strip())
    if len(names) > 6:
        return ", ".join(names[:6]) + ", et al."
    if len(names) == 1:
        return names[0]
    return ", ".join(names[:-1]) + (", and " if len(names) > 2 else " and ") + names[-1]


def fmt_ref(e):
    """IEEE reference string; *...* marks italics."""
    t, au, ti, yr = e.get("ENTRYTYPE"), fmt_authors(e["author"]), tex_to_text(e["title"]), e.get("year", "")
    if t == "article":
        return f'{au}, "{ti}," *{tex_to_text(e["journal"])}*, vol. {e["volume"]}, pp. {tex_to_text(e["pages"])}, {yr}.'
    if t == "inproceedings":
        s = f'{au}, "{ti}," in *{tex_to_text(e["booktitle"])}*'
        if "volume" in e:
            s += f", vol. {e['volume']}"
        s += f", {yr}"
        if "pages" in e:
            s += f', pp. {tex_to_text(e["pages"])}'
        return s + "."
    return f'{au}, "{ti}," arXiv preprint arXiv:{e.get("eprint", "")}, {yr}.'


def emit_refs(doc, N, entries):
    by_key = {e["ID"]: e for e in entries}
    for i, key in enumerate(N.cites, 1):
        if key not in by_key:
            WARN.append(f"missing bib {key}")
            continue
        p = new_para(doc, indent=False, size=8)
        pf = p.paragraph_format
        pf.left_indent, pf.first_line_indent = Inches(0.25), Inches(-0.25)
        add_tab(p, Inches(0.25), WD_TAB_ALIGNMENT.LEFT)
        run(p, f"[{i}]\t", size=8)
        for j, chunk in enumerate(fmt_ref(by_key[key]).split("*")):
            if chunk:
                run(p, chunk, italic=(j % 2 == 1), size=8)


# ----------------------------------------------------------------- main
def main():
    main_tex = strip_comments(open(os.path.join(ROOT, "main.tex"), encoding="utf-8").read())
    body = main_tex[main_tex.index("\\maketitle") + len("\\maketitle"):main_tex.index("\\bibliographystyle")]
    body = re.sub(r"\\input\{([^}]*)\}",
                  lambda m: strip_comments(open(os.path.join(ROOT, m.group(1) + ".tex"), encoding="utf-8").read()),
                  body)
    abstract = re.search(r"\\begin\{abstract\}(.*?)\\end\{abstract\}", body, re.S).group(1)
    keywords = re.search(r"\\begin\{IEEEkeywords\}(.*?)\\end\{IEEEkeywords\}", body, re.S).group(1)
    body = re.sub(r"\\begin\{abstract\}.*?\\end\{abstract\}|\\begin\{IEEEkeywords\}.*?\\end\{IEEEkeywords\}",
                  "", body, flags=re.S)
    N = Numbering(body)
    with open(os.path.join(ROOT, "refs.bib"), encoding="utf-8") as f:
        entries = bibtexparser.load(f).entries

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Times New Roman", Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.paragraph_format.space_after, st.paragraph_format.line_spacing = Pt(0), 1.0
    doc.styles["List Bullet"].font.name, doc.styles["List Bullet"].font.size = "Times New Roman", Pt(10)
    s0 = doc.sections[0]
    page_setup(s0)
    set_cols(s0, 1)

    # title block (single column)
    title = find_arg(main_tex, "title")[0]
    p = new_para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, after=14)
    for j, line in enumerate(re.split(r"\\\\", title)):
        if j:
            p.add_run().add_break(WD_BREAK.LINE)
        run(p, norm_ws(line), size=24)
    authors = re.split(r"\\and", find_arg(main_tex, "author")[0])
    t = doc.add_table(rows=1, cols=len(authors))
    t.alignment = 1
    for cell, a in zip(t.rows[0].cells, authors):
        name, aff = find_arg(a, "IEEEauthorblockN")[0], find_arg(a, "IEEEauthorblockA")[0]
        cp = cell.paragraphs[0]
        cp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp._size = 11
        emit_inline(cp, norm_ws(name), N)
        for line in re.split(r"\\\\", aff):
            q = cell.add_paragraph()
            q.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            q._size = 10
            emit_inline(q, norm_ws(line), N)
        set_cell_borders(cell, top=0, bottom=0, left=0, right=0)
    new_para(doc, indent=False, after=6)

    # body (two columns)
    s1 = doc.add_section(WD_SECTION.CONTINUOUS)
    page_setup(s1)
    set_cols(s1, 2)
    p = new_para(doc, indent=False, after=4, size=9)
    run(p, "Abstract\u2014", bold=True, italic=True, size=9)
    emit_inline(p, norm_ws(abstract), N, bold=True)
    p = new_para(doc, indent=False, after=6, size=9)
    run(p, "Index Terms\u2014", bold=True, italic=True, size=9)
    run(p, norm_ws(keywords), bold=True, size=9)

    B = Builder(doc, N)
    B.body(body)
    B.heading("References", None)
    emit_refs(doc, N, entries)
    s_end = doc.add_section(WD_SECTION.CONTINUOUS)  # trailing break makes Word balance the last page's columns
    page_setup(s_end)
    set_cols(s_end, 2)
    doc.save(OUT)
    print("wrote", OUT)
    print("cites:", N.cites)
    print("labels:", N.labels)
    for w in WARN:
        print("WARN", w)


if __name__ == "__main__":
    main()
